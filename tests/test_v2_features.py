"""Attachments and HTML email, forwarded emails, customer price lists, shipping API, n8n sending,
database upgrades and serving the web app."""

import base64
import email
import io
import json
from decimal import Decimal as D
from email import policy
from email.message import EmailMessage
from types import SimpleNamespace

import httpx
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from reportlab.pdfgen import canvas
from sqlalchemy import text

import app.models  # noqa: F401  (registers tables on Base.metadata)
from app.db import Base, make_engine
from app.documents import find_forwarded_sender, html_to_text, read_email
from app.extraction import OpenAIExtractor
from app.main import create_app
from app.migrate import ensure_schema
from app.schemas import Attachment, ExtractedQuoteRequest, InboundEmail
from app.shipping import HttpShipping, RulesShipping
from tests.conftest import AUTH, WEBHOOK, get_quote, ingest, lines_by_sku

JORDAN = "jordan.lee@brightline-construction.com"


def b64(data: bytes) -> str:
    return base64.b64encode(data).decode()


def xlsx_bytes(rows) -> bytes:
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def docx_bytes(rows) -> bytes:
    document = Document()
    document.add_paragraph("Please quote the items below.")
    table = document.add_table(rows=0, cols=len(rows[0]))
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def scanned_pdf_bytes() -> bytes:
    """A PDF with no text layer, like a scan."""
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.rect(50, 50, 200, 100, fill=1)
    pdf.showPage()
    pdf.save()
    return buffer.getvalue()


# ------------------------------------------------------------------ reading emails and attachments

def test_html_to_text_keeps_table_rows():
    html = (
        "<html><head><style>p{color:red}</style></head><body><p>Hi&nbsp;team</p>"
        "<table><tr><td>SKU</td><td>Qty</td></tr><tr><td>FS-HN-M8</td><td>40</td></tr></table>"
        "<div>Thanks<br>Jordan</div></body></html>"
    )
    assert html_to_text(html) == "Hi team\nSKU | Qty\nFS-HN-M8 | 40\nThanks\nJordan"


def test_read_email_reports_every_attachment():
    original = EmailMessage()
    original["From"] = f"Jordan Lee <{JORDAN}>"
    original["Subject"] = "RFQ"
    original.set_content("- 5 x FS-WA-M8")
    inbound = InboundEmail(
        message_id="m",
        from_email="sam@acme-supply.example",
        body_html="<p>See&nbsp;below</p><table><tr><td>FS-HN-M8</td><td>40</td></tr></table>",
        attachments=[
            Attachment(filename="original.eml", content_type="message/rfc822", content_base64=b64(bytes(original))),
            Attachment(filename="old.xls", content_type="application/vnd.ms-excel", content_base64=b64(b"legacy")),
            Attachment(filename="broken.pdf", content_type="application/pdf", content_base64="not base64!"),
            Attachment(filename="scan.pdf", content_type="application/pdf", content_base64=b64(scanned_pdf_bytes())),
        ],
    )
    content = read_email(inbound)
    assert content.body == "See below\nFS-HN-M8 | 40"
    assert f"From: Jordan Lee <{JORDAN}>" in content.attachments_text and "- 5 x FS-WA-M8" in content.attachments_text
    assert [entry["method"] for entry in content.report] == ["text", "unsupported", "error", "ai"]
    assert [f.filename for f in content.ai_files] == ["scan.pdf"]


def test_find_forwarded_sender_skips_own_domain():
    text_block = (
        "From: Sam Ortiz <sam@acme-supply.example>\nSent: Monday\n\n"
        f"*From:* Lee, Jordan [mailto:{JORDAN}]\nSubject: RFQ\n"
    )
    assert find_forwarded_sender(text_block, frozenset({"acme-supply.example"})) == ("Lee, Jordan", JORDAN)
    assert find_forwarded_sender("no headers here", frozenset()) is None


def test_excel_attachment_becomes_quote_lines(client):
    workbook = xlsx_bytes([["SKU", "Description", "Qty"], ["FS-HN-M8", "M8 nuts", 40], ["FS-WA-M8", "M8 washers", 25]])
    result = ingest(client, {
        "message_id": "<xlsx-1@brightline>", "from_email": JORDAN, "subject": "RFQ attached",
        "body_text": "Hi, our list is attached.",
        "attachments": [{"filename": "rfq.xlsx", "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "content_base64": b64(workbook)}],
    })
    quote = get_quote(client, result["quote_id"])
    assert {line["sku"]: D(line["quantity"]) for line in quote["lines"]} == {"FS-HN-M8": D(40), "FS-WA-M8": D(25)}
    assert quote["attachments"][0]["method"] == "text"


def test_word_attachment_becomes_quote_lines(client):
    result = ingest(client, {
        "message_id": "<docx-1@brightline>", "from_email": JORDAN, "subject": "RFQ",
        "attachments": [{"filename": "rfq.docx", "content_type": "application/octet-stream", "content_base64": b64(docx_bytes([["SKU", "Qty"], ["FS-HB-M8-50", "20"]]))}],
    })
    [line] = get_quote(client, result["quote_id"])["lines"]
    assert (line["sku"], D(line["quantity"]), line["match_method"]) == ("FS-HB-M8-50", D(20), "sku")


def test_html_only_email_is_read(client):
    result = ingest(client, {
        "message_id": "<html-1@brightline>", "from_email": JORDAN, "subject": "RFQ",
        "body_html": "<p>Please quote:</p><ul><li>10 x FS-WA-M8</li></ul>",
    })
    quote = get_quote(client, result["quote_id"])
    assert quote["source_body"] == "Please quote:\n10 x FS-WA-M8"
    assert [line["sku"] for line in quote["lines"]] == ["FS-WA-M8"]


def test_scanned_pdf_is_flagged_when_ai_reading_is_off(client):
    result = ingest(client, {
        "message_id": "<scan-1@brightline>", "from_email": JORDAN, "subject": "RFQ",
        "body_text": "Please see the attached scan.",
        "attachments": [{"filename": "scan.pdf", "content_type": "application/pdf", "content_base64": b64(scanned_pdf_bytes())}],
    })
    assert result["status"] == "pending_approval"
    quote = get_quote(client, result["quote_id"])
    assert quote["lines"] == [] and quote["attachments"][0]["method"] == "ai"
    assert "scan.pdf" in quote["internal_notes"]


def test_openai_extractor_sends_scans_and_photos_to_the_model():
    captured = {}

    class FakeResponses:
        def parse(self, **kwargs):
            captured.update(kwargs)
            empty = ExtractedQuoteRequest(is_quote_request=True, customer_name=None, company=None, customer_email=None,
                                          shipping_address=None, requested_delivery=None, items=[], notes=None)
            return SimpleNamespace(output_parsed=empty)

    extractor = OpenAIExtractor("sk-test", "test-model", client=SimpleNamespace(responses=FakeResponses()))
    inbound = InboundEmail(message_id="m", from_email="a@b.com", body_text="see attached", attachments=[
        Attachment(filename="scan.pdf", content_type="application/pdf", content_base64=b64(scanned_pdf_bytes())),
        Attachment(filename="photo.png", content_type="image/png", content_base64=b64(b"\x89PNG\r\n\x1a\n" + b"0" * 20)),
    ])
    extractor.extract(inbound, read_email(inbound))
    parts = captured["input"][1]["content"]
    assert [part["type"] for part in parts] == ["input_text", "input_file", "input_image"]
    assert parts[1]["file_data"].startswith("data:application/pdf;base64,")
    assert parts[2]["image_url"].startswith("data:image/png;base64,")
    assert "scan.pdf" in parts[0]["text"]


# ------------------------------------------------------------------ forwarded emails

FORWARDED_BODY = (
    "Can we quote this for Jordan?\n\n"
    "---------- Forwarded message ---------\n"
    f"From: Jordan Lee <{JORDAN}>\n"
    "Date: Mon, 14 Sep 2026 at 09:12\n"
    "Subject: Quote request\n"
    "To: Sam Ortiz <sam@acme-supply.example>\n\n"
    "- 25 x hard hats (white)\n"
    "- 40 boxes M8 hex nuts\n"
)


def test_forwarded_email_quotes_the_original_customer_and_copies_the_colleague(client, settings):
    result = ingest(client, {"message_id": "<fwd-1@acme>", "from_email": "sam@acme-supply.example", "from_name": "Sam Ortiz",
                             "subject": "Fwd: Quote request", "body_text": FORWARDED_BODY})
    quote = get_quote(client, result["quote_id"])
    assert quote["customer"]["email"] == JORDAN and quote["forwarded_by"] == "sam@acme-supply.example"
    assert lines_by_sku(quote)["SF-HH-WHT"]["price_source"] == "customer_price"

    sent = client.post(f"/api/quotes/{quote['id']}/approve", json={"approver": "Sam"}, headers=AUTH).json()
    assert sent["status"] == "sent"
    [eml] = list(settings.outbox_dir.glob("*.eml"))
    message = email.message_from_bytes(eml.read_bytes(), policy=policy.default)
    assert (message["To"], message["Cc"]) == (JORDAN, "sam@acme-supply.example")
    assert message["In-Reply-To"] is None
    assert message["Subject"] == f"Quotation {sent['number']} - Quote request"


def test_forward_without_customer_address_must_be_assigned_before_approval(client):
    result = ingest(client, {"message_id": "<fwd-2@acme>", "from_email": "sam@acme-supply.example",
                             "subject": "Fwd: phone order", "body_text": "Customer called, please quote:\n- 10 x FS-WA-M8"})
    quote_id = result["quote_id"]
    quote = get_quote(client, quote_id)
    assert quote["customer"] is None and "could not be found" in quote["internal_notes"]
    assert client.post(f"/api/quotes/{quote_id}/approve", json={"approver": "Sam"}, headers=AUTH).status_code == 409
    assert client.patch(f"/api/quotes/{quote_id}", json={"actor": "Sam", "customer_email": "sam@acme-supply.example"}, headers=AUTH).status_code == 409

    updated = client.patch(f"/api/quotes/{quote_id}", json={"actor": "Sam", "customer_email": "priya@northwindfacilities.com"}, headers=AUTH).json()
    assert updated["customer"]["email"] == "priya@northwindfacilities.com"
    assert D(updated["lines"][0]["discount_pct"]) == D("10") and D(updated["tax_rate"]) == D("0.0625")
    assert "could not be found" not in (updated["internal_notes"] or "")
    assert client.post(f"/api/quotes/{quote_id}/approve", json={"approver": "Sam"}, headers=AUTH).status_code == 200


# ------------------------------------------------------------------ customer price lists and shipping

def test_customer_price_list_api(client):
    customers = client.get("/api/customers", headers=AUTH).json()
    priya = next(c for c in customers if c["email"] == "priya@northwindfacilities.com")
    assert priya["price_count"] == 1

    created = client.put(f"/api/customers/{priya['id']}/prices", json={"sku": "SF-HH-WHT", "unit_price": "10.75"}, headers=AUTH).json()
    assert created["product_name"].startswith("Safety Helmet") and D(created["list_price"]) == D("14.90")
    result = ingest(client, {"message_id": "<priya-1@northwind>", "from_email": "priya@northwindfacilities.com", "subject": "RFQ", "body_text": "- 4 x SF-HH-WHT"})
    [line] = get_quote(client, result["quote_id"])["lines"]
    assert (D(line["unit_price"]), line["price_source"], D(line["discount_pct"])) == (D("10.75"), "customer_price", D("0"))

    assert client.put(f"/api/customers/{priya['id']}/prices", json={"sku": "NOPE", "unit_price": "1"}, headers=AUTH).status_code == 409
    assert client.delete(f"/api/customers/{priya['id']}/prices/{created['id']}", headers=AUTH).status_code == 204
    assert len(client.get(f"/api/customers/{priya['id']}/prices", headers=AUTH).json()) == 1

    updated = client.patch(f"/api/customers/{priya['id']}", json={"payment_terms": "Net 60", "is_verified": True}, headers=AUTH).json()
    assert updated["payment_terms"] == "Net 60"
    assert client.patch(f"/api/customers/{priya['id']}", json={"tax_region": "XX-NOPE"}, headers=AUTH).status_code == 409


def test_shipping_api_price_is_used_and_falls_back_to_rules(settings):
    def rates(request):
        body = json.loads(request.content)
        assert D(body["weight_kg"]) == D("5.4") and body["items"][0]["sku"] == "EL-PP-24-C6"
        return httpx.Response(200, json={"amount": "22.40", "service": "Ground", "transit_days": 3})

    rules = RulesShipping(D("1.50"), D("2500"), D("15"))
    cases = [
        (httpx.MockTransport(rates), ("22.40", "api", "Ground", 3)),
        (httpx.MockTransport(lambda request: httpx.Response(500)), ("15.00", "rules_fallback", "Standard", None)),
    ]
    for transport, expected in cases:
        shipping = HttpShipping("http://ship.test", "key", rules, client=httpx.Client(transport=transport))
        with TestClient(create_app(settings, shipping=shipping)) as client:
            result = ingest(client, {"message_id": f"<ship-{expected[1]}@x>", "from_email": "alex@gmail.com", "subject": "quote", "body_text": "- 3 x 24 port patch panel"})
            quote = get_quote(client, result["quote_id"])
            assert (quote["shipping_total"], quote["shipping_source"], quote["shipping_service"], quote["shipping_transit_days"]) == expected


# ------------------------------------------------------------------ sending through n8n

def test_n8n_send_mode_claims_once_and_reports_back(settings, sample_email):
    handoffs = []

    def n8n(request):
        handoffs.append(json.loads(request.content))
        return httpx.Response(200, json={"message": "Workflow was started"})

    n8n_settings = settings.model_copy(
        update={"send_mode": "n8n", "n8n_quote_approved_webhook_url": "http://n8n.test/webhook/quote-approved", "allow_sample_sends": True}
    )
    with TestClient(create_app(n8n_settings, http=httpx.Client(transport=httpx.MockTransport(n8n)))) as client:
        quote_id = ingest(client, sample_email)["quote_id"]
        approved = client.post(f"/api/quotes/{quote_id}/approve", json={"approver": "Sam", "confirm_flagged_lines": True}, headers=AUTH).json()
        assert approved["status"] == "approved" and approved["events"][-1]["event"] == "handed_to_n8n"
        assert handoffs[0]["event"] == "quote.approved" and handoffs[0]["quote_id"] == quote_id

        claim_url = f"/api/integrations/quotes/{quote_id}/claim"
        assert client.post(claim_url).status_code == 401  # needs the webhook secret
        claim = client.post(claim_url, headers=WEBHOOK).json()
        assert claim["to"] == JORDAN and claim["cc"] == [] and claim["subject"] == "Re: Quote request - Oakland site"
        assert base64.b64decode(claim["attachment_base64"]).startswith(b"%PDF")

        assert client.post(claim_url, headers=WEBHOOK).status_code == 409  # a second run can't send it twice
        sent_url = f"/api/integrations/quotes/{quote_id}/sent"
        assert client.post(sent_url, json={"claim_token": "wrong"}, headers=WEBHOOK).status_code == 409

        failed = client.post(f"/api/integrations/quotes/{quote_id}/failed", json={"claim_token": claim["claim_token"], "error": "Gmail quota"}, headers=WEBHOOK)
        assert failed.json()["status"] == "approved"

        retry = client.post(claim_url, headers=WEBHOOK).json()
        done = client.post(sent_url, json={"claim_token": retry["claim_token"], "provider_message_id": "gmail-123"}, headers=WEBHOOK).json()
        assert done["status"] == "sent"
        assert get_quote(client, quote_id)["events"][-1]["detail"]["via"] == "n8n"


def test_n8n_handoff_failure_is_reported(settings, sample_email):
    n8n_settings = settings.model_copy(
        update={"send_mode": "n8n", "n8n_quote_approved_webhook_url": "http://n8n.test/down", "allow_sample_sends": True}
    )
    down = httpx.Client(transport=httpx.MockTransport(lambda request: httpx.Response(503)))
    with TestClient(create_app(n8n_settings, http=down)) as client:
        quote_id = ingest(client, sample_email)["quote_id"]
        response = client.post(f"/api/quotes/{quote_id}/approve", json={"approver": "Sam", "confirm_flagged_lines": True}, headers=AUTH)
        assert response.status_code == 502
        quote = get_quote(client, quote_id)
        assert quote["status"] == "approved" and quote["events"][-1]["event"] == "handoff_failed"


# ------------------------------------------------------------------ upgrades and web app

def test_old_database_gets_new_columns(tmp_path):
    engine = make_engine(f"sqlite:///{tmp_path / 'old.db'}")
    Base.metadata.create_all(engine)
    with engine.begin() as conn:
        conn.execute(text("ALTER TABLE quote_lines DROP COLUMN price_source"))
        conn.execute(text("ALTER TABLE quotes DROP COLUMN forwarded_by"))
        conn.execute(text("DROP TABLE customer_prices"))
    assert sorted(ensure_schema(engine, Base.metadata)) == ["quote_lines.price_source", "quotes.forwarded_by"]
    assert ensure_schema(engine, Base.metadata) == []


def test_web_app_is_served_once_built(settings, tmp_path):
    with TestClient(create_app(settings)) as client:
        page = client.get("/")
        assert page.status_code == 503 and "npm run build" in page.text

    web = tmp_path / "web-out"
    (web / "customers" / "__next.customers").mkdir(parents=True)
    (web / "index.html").write_text("<h1>Quote approvals</h1>", encoding="utf-8")
    (web / "customers" / "index.html").write_text("<h1>Customers</h1>", encoding="utf-8")
    (web / "customers" / "__next.customers" / "__PAGE__.txt").write_text("prefetch-data", encoding="utf-8")
    with TestClient(create_app(settings.model_copy(update={"web_dist_dir": web}))) as client:
        assert "Quote approvals" in client.get("/").text
        assert "Customers" in client.get("/customers/").text
        # Next.js asks for the dotted name; the export stores it in a folder.
        prefetch = client.get("/customers/__next.customers.__PAGE__.txt?_rsc=abc")
        assert prefetch.status_code == 200 and prefetch.text == "prefetch-data"
        assert client.get("/customers/__next.nothing.__PAGE__.txt").status_code == 404
        assert client.get("/health").json() == {"status": "ok"}
        assert client.get("/api/config", headers=AUTH).json()["internal_email_domains"] == ["acme-supply.example"]
