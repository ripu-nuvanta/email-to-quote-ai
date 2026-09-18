"""Read everything a customer might send: plain or HTML email bodies, forwarded messages, and
attachments (PDF, Excel, Word, CSV/TXT, HTML, attached .eml messages, scanned PDFs and photos).

Text-based files are converted to text here. Scanned PDFs and images can't be read locally, so they are
handed to the AI model as files (see OpenAIExtractor). Every attachment gets a report entry so the
salesperson can see what was, and wasn't, read.
"""

from __future__ import annotations

import base64
import binascii
import io
import logging
import re
from dataclasses import dataclass, field
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
from pathlib import PurePath

from pypdf import PdfReader

from .schemas import InboundEmail

log = logging.getLogger(__name__)

MAX_TEXT_CHARS = 20_000
MAX_SHEET_ROWS = 500
MAX_AI_FILES = 5
SCANNED_PDF_MIN_CHARS = 40

XLSX_TYPES = frozenset(
    {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/vnd.ms-excel.sheet.macroenabled.12"}
)
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp", ".gif": "image/gif"}


@dataclass(frozen=True)
class AIFile:
    """A file only an AI model can read (a scanned PDF or an image)."""

    filename: str
    mime_type: str
    data: bytes

    @property
    def data_url(self) -> str:
        return f"data:{self.mime_type};base64,{base64.b64encode(self.data).decode()}"


@dataclass
class EmailContent:
    body: str
    attachments_text: str = ""
    ai_files: list[AIFile] = field(default_factory=list)
    report: list[dict] = field(default_factory=list)  # per attachment: filename, content_type, size_bytes, method, note

    @property
    def full_text(self) -> str:
        return f"{self.body}\n{self.attachments_text}"


# --------------------------------------------------------------------------- HTML

_BLOCK_TAGS = frozenset(
    {"p", "div", "br", "tr", "li", "ul", "ol", "table", "section", "article", "blockquote", "hr", "pre", "h1", "h2", "h3", "h4", "h5", "h6"}
)
_SKIP_TAGS = frozenset({"script", "style", "head", "title"})


class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in _SKIP_TAGS:
            self._skip += 1
        elif tag in ("td", "th"):
            self.parts.append(" | ")
        elif tag in _BLOCK_TAGS:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in _SKIP_TAGS:
            self._skip = max(0, self._skip - 1)
        elif tag in _BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def html_to_text(html: str) -> str:
    parser = _HTMLText()
    parser.feed(html)
    parser.close()
    lines = (re.sub(r"[ \t ]+", " ", line).strip().strip("|").strip() for line in "".join(parser.parts).splitlines())
    return "\n".join(line for line in lines if line)


def email_body(email: InboundEmail) -> str:
    if email.body_text.strip():
        return email.body_text
    return html_to_text(email.body_html or "")


# --------------------------------------------------------------------------- attachments

def read_email(email: InboundEmail, max_attachment_bytes: int = 15 * 1024 * 1024) -> EmailContent:
    content = EmailContent(body=email_body(email))
    texts: list[str] = []
    for att in email.attachments:
        entry = {"filename": att.filename, "content_type": att.content_type, "size_bytes": None, "method": "unsupported", "note": None}
        content.report.append(entry)
        try:
            data = base64.b64decode(att.content_base64, validate=True)
        except (binascii.Error, ValueError):
            entry.update(method="error", note="File data was not valid base64")
            continue
        entry["size_bytes"] = len(data)
        if len(data) > max_attachment_bytes:
            entry["note"] = f"Larger than {max_attachment_bytes // (1024 * 1024)} MB"
            continue
        try:
            kind, result = _read_attachment(att.filename, att.content_type, data)
        except Exception as exc:  # a broken file must not block the rest of the email
            log.warning("Could not read attachment %s", att.filename, exc_info=True)
            entry.update(method="error", note=f"Could not open the file ({type(exc).__name__})")
            continue
        if kind == "text":
            entry["method"] = "text"
            texts.append(f"--- Attachment: {att.filename} ---\n{result[:MAX_TEXT_CHARS]}")
        elif kind == "ai":
            if len(content.ai_files) >= MAX_AI_FILES:
                entry["note"] = f"More than {MAX_AI_FILES} scanned files/images; skipped"
                continue
            entry["method"] = "ai"
            content.ai_files.append(result)
        else:
            entry["note"] = result
    content.attachments_text = "\n\n".join(texts)
    return content


def _read_attachment(filename: str, content_type: str, data: bytes) -> tuple[str, str | AIFile]:
    ext = PurePath(filename.lower()).suffix
    ctype = (content_type or "").lower().split(";")[0].strip()

    if ext == ".pdf" or ctype == "application/pdf":
        text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
        if len(text.strip()) >= SCANNED_PDF_MIN_CHARS:
            return "text", text
        return "ai", AIFile(filename, "application/pdf", data)  # no text layer: a scan
    if ext in IMAGE_TYPES or ctype.startswith("image/"):
        mime = IMAGE_TYPES.get(ext, ctype)
        if mime not in IMAGE_TYPES.values():
            return "unsupported", f"Image type {mime} is not supported"
        return "ai", AIFile(filename, mime, data)
    if ext in (".xlsx", ".xlsm") or ctype in XLSX_TYPES:
        return "text", _excel_text(data)
    if ext == ".docx" or ctype == DOCX_TYPE:
        return "text", _word_text(data)
    if ext in (".html", ".htm") or ctype == "text/html":
        return "text", html_to_text(data.decode("utf-8", errors="replace"))
    if ext == ".eml" or ctype == "message/rfc822":
        return "text", _eml_text(data)
    if ext in (".txt", ".csv", ".tsv", ".md") or ctype.startswith("text/"):
        return "text", data.decode("utf-8", errors="replace")
    if ext in (".xls", ".doc"):
        return "unsupported", "Old Office format: ask for .xlsx / .docx or PDF"
    return "unsupported", "File type not supported"


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _excel_text(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = []
    try:
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cells = [_cell(v) for v in row]
                if any(cells):
                    rows.append(" | ".join(c for c in cells if c))
                if len(rows) >= MAX_SHEET_ROWS:
                    break
            if rows:
                sheets.append(f"[Sheet: {sheet.title}]\n" + "\n".join(rows))
    finally:
        workbook.close()
    return "\n\n".join(sheets)


def _word_text(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(c for c in cells if c))
    return "\n".join(parts)


def _eml_text(data: bytes) -> str:
    message = BytesParser(policy=policy.default).parsebytes(data)
    body = message.get_body(preferencelist=("plain", "html"))
    text = ""
    if body is not None:
        text = body.get_content()
        if body.get_content_type() == "text/html":
            text = html_to_text(text)
    return f"From: {message.get('From', '')}\nSubject: {message.get('Subject', '')}\n\n{text}"


# --------------------------------------------------------------------------- forwarded emails

_FROM_LINE = re.compile(r"^[>\s*]*From:\s*\**\s*(.+)$", re.I | re.M)
_ADDRESS = re.compile(r"[\w.+'-]+@[\w-]+(?:\.[\w-]+)+")


def find_forwarded_sender(text: str, internal_domains: frozenset[str]) -> tuple[str | None, str] | None:
    """Return (name, email) from the first 'From:' line in forwarded content that isn't one of our addresses."""
    for match in _FROM_LINE.finditer(text):
        value = match.group(1)
        address = _ADDRESS.search(value)
        if not address:
            continue
        email = address.group(0).lower()
        if email.rpartition("@")[2] in internal_domains:
            continue
        name = re.sub(r"\[?mailto:\s*$", "", value[: address.start()], flags=re.I).strip(" \"'<[(*")
        return (name or None), email
    return None
